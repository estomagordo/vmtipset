#!/usr/bin/env python3
"""
Extract structure from a .xlsx for rebuilding the UX in a web app.

- Sheet names, used range, merged cells
- Every non-empty cell: value vs formula, data types

Usage:
  python -m venv .venv && source .venv/bin/activate
  pip install -r requirements.txt
  python tools/parse_workbook.py "VM-tipset-2026-Mall version 1.3.xlsx"
  python tools/parse_workbook.py workbook.xlsx --json output.json
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any


def _cell_payload(cell) -> dict[str, Any]:
    """openpyxl cell -> serializable dict."""
    coord = cell.coordinate
    dt = cell.data_type
    number_format = cell.number_format or ""

    # Formula cells: data_type == 'f'. value is the formula string (e.g. "=A1+B1").
    if dt == "f":
        return {
            "address": coord,
            "kind": "formula",
            "formula": cell.value,
            "number_format": number_format,
        }

    v = cell.value
    if v is None:
        return {}

    kind = "empty"
    if dt == "n":
        kind = "number"
    elif dt == "s":
        kind = "string"
    elif dt == "b":
        kind = "bool"
    elif dt == "d":
        kind = "date"
    elif dt == "e":
        kind = "error"
    else:
        kind = dt or "other"

    out: dict[str, Any] = {
        "address": coord,
        "kind": kind,
        "value": v,
    }
    if number_format and number_format != "General":
        out["number_format"] = number_format
    return out


def parse_workbook(path: Path, *, data_only: bool) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=data_only, read_only=False)

    sheets: list[dict[str, Any]] = []
    for name in wb.sheetnames:
        ws = wb[name]
        merged = [str(rng) for rng in ws.merged_cells.ranges]

        cells: list[dict[str, Any]] = []
        # iter_rows with values_only=False gives real Cell objects (formulas when data_only=False).
        for row in ws.iter_rows():
            for cell in row:
                payload = _cell_payload(cell)
                if payload:
                    cells.append(payload)

        sheets.append(
            {
                "name": name,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "merged_ranges": merged,
                "cells": cells,
            }
        )

    wb.close()
    return {
        "workbook": path.name,
        "data_only": data_only,
        "sheet_count": len(sheets),
        "sheets": sheets,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Parse .xlsx to JSON (values + formulas).")
    parser.add_argument(
        "workbook",
        type=Path,
        nargs="?",
        default=Path("VM-tipset-2026-Mall version 1.3.xlsx"),
        help="Path to .xlsx (default: VM-tipset-2026-Mall version 1.3.xlsx)",
    )
    parser.add_argument(
        "--json",
        type=Path,
        metavar="FILE",
        help="Write full dump to FILE (UTF-8)",
    )
    parser.add_argument(
        "--data-only",
        action="store_true",
        help="Load cached values only (formulas become computed values or None if no cache)",
    )
    parser.add_argument(
        "--summary",
        action="store_true",
        help="Print per-sheet counts only (no cell listing on stdout)",
    )
    args = parser.parse_args()

    wb_path: Path = args.workbook
    if not wb_path.is_file():
        print(f"error: not a file: {wb_path}", file=sys.stderr)
        return 1

    data = parse_workbook(wb_path, data_only=args.data_only)

    if args.json:
        args.json.write_text(
            json.dumps(data, ensure_ascii=False, indent=2, default=str) + "\n",
            encoding="utf-8",
        )
        print(f"wrote {args.json}")

    if args.summary:
        for sh in data["sheets"]:
            formulas = sum(1 for c in sh["cells"] if c.get("kind") == "formula")
            values = len(sh["cells"]) - formulas
            print(
                f'{sh["name"]}: rows≤{sh["max_row"]} cols≤{sh["max_column"]} '
                f"cells={len(sh['cells'])} (values≈{values}, formulas≈{formulas}) "
                f'merged={len(sh["merged_ranges"])}'
            )
    elif not args.json:
        # Default: pretty JSON to stdout (can be large).
        json.dump(data, sys.stdout, ensure_ascii=False, indent=2, default=str)
        sys.stdout.write("\n")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
